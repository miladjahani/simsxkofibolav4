import os, json, re, shutil, tempfile, subprocess, secrets
from datetime import datetime, date
from pathlib import Path
from typing import Optional, Any
from fastapi import FastAPI, Depends, HTTPException, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sqlalchemy import create_engine, String, Integer, Float, DateTime, Text, ForeignKey, select
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, Session, sessionmaker
from passlib.context import CryptContext
from jose import jwt
from .native_engine import NativeWorkbook, FormulaError
from .process_engine import SXEWProcessEngine, ProcessError
from .optimizer import optimize, OptimizationError
import csv, io
import openpyxl

BASE=Path(__file__).resolve().parent
DATA=BASE/"data"
MODEL=DATA/"model.xlsx"
META=DATA/"workbook_meta.json"
SECRET=os.getenv("SECRET_KEY","dev-only-change-me")
DBURL=os.getenv("DATABASE_URL","sqlite:///./sxew.db")
connect_args={"check_same_thread":False} if DBURL.startswith("sqlite") else {}
engine=create_engine(DBURL,connect_args=connect_args)
SessionLocal=sessionmaker(bind=engine,autocommit=False,autoflush=False)
pwd=CryptContext(schemes=["bcrypt"],deprecated="auto")
app=FastAPI(title="SX-EW Digital Simulator", version="1.0.0")
app.mount("/static",StaticFiles(directory=str(BASE/"static")),name="static")

class Base(DeclarativeBase): pass
class User(Base):
    __tablename__="users"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    email:Mapped[str]=mapped_column(String(200),unique=True,index=True)
    password_hash:Mapped[str]=mapped_column(String(200))
    created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow)
class Scenario(Base):
    __tablename__="scenarios"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    user_id:Mapped[int]=mapped_column(ForeignKey("users.id"))
    name:Mapped[str]=mapped_column(String(200))
    sheet:Mapped[str]=mapped_column(String(200))
    config:Mapped[str]=mapped_column(Text)
    created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow)
class Timeseries(Base):
    __tablename__="timeseries"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    user_id:Mapped[int]=mapped_column(ForeignKey("users.id"))
    ts:Mapped[datetime]=mapped_column(DateTime,index=True)
    metric:Mapped[str]=mapped_column(String(100),index=True)
    value:Mapped[float]=mapped_column(Float)
    period:Mapped[str]=mapped_column(String(20),default="daily")
Base.metadata.create_all(engine)

def db():
    s=SessionLocal()
    try: yield s
    finally: s.close()
def token_for(u):
    return jwt.encode({"sub":str(u.id),"exp":datetime.utcnow().timestamp()+86400*7},SECRET,algorithm="HS256")
def current_user(req:Request,s:Session=Depends(db)):
    h=req.headers.get("authorization","")
    if not h.startswith("Bearer "): raise HTTPException(401,"Login required")
    try: uid=int(jwt.decode(h[7:],SECRET,algorithms=["HS256"])["sub"])
    except Exception: raise HTTPException(401,"Invalid token")
    u=s.get(User,uid)
    if not u: raise HTTPException(401,"Invalid user")
    return u

class Auth(BaseModel): email:str; password:str
class Calc(BaseModel): sheet:str; inputs:dict[str,Any]={}
class ScenarioIn(BaseModel): name:str; sheet:str; config:dict[str,Any]
class DataIn(BaseModel): ts:datetime; metric:str; value:float; period:str="daily"

@app.get("/",response_class=HTMLResponse)
def home():
    return (BASE/"templates/index.html").read_text(encoding="utf8")

@app.post("/api/auth/register")
def register(x:Auth,s:Session=Depends(db)):
    if s.scalar(select(User).where(User.email==x.email)): raise HTTPException(400,"Email already exists")
    u=User(email=x.email.lower().strip(),password_hash=pwd.hash(x.password)); s.add(u); s.commit(); s.refresh(u)
    return {"token":token_for(u),"email":u.email}
@app.post("/api/auth/login")
def login(x:Auth,s:Session=Depends(db)):
    u=s.scalar(select(User).where(User.email==x.email.lower().strip()))
    if not u or not pwd.verify(x.password,u.password_hash): raise HTTPException(401,"Wrong email or password")
    return {"token":token_for(u),"email":u.email}

@app.get("/api/workbook")
def workbook():
    return json.loads(META.read_text(encoding="utf8"))

@app.get("/api/workbook/{sheet}")
def sheet_detail(sheet:str):
    wb=openpyxl.load_workbook(MODEL,data_only=False,read_only=True)
    if sheet not in wb.sheetnames: raise HTTPException(404,"Sheet not found")
    ws=wb[sheet]
    rows=[]
    for r in ws.iter_rows():
        vals=[c.value for c in r]
        if any(v is not None for v in vals):
            rows.append({"row":r[0].row,"cells":[{"cell":c.coordinate,"value":c.value} for c in r if c.value is not None]})
    return {"sheet":sheet,"rows":rows}

@app.post("/api/calculate")
def calculate(x:Calc,user=Depends(current_user)):
    if x.sheet not in openpyxl.load_workbook(MODEL,read_only=True).sheetnames: raise HTTPException(404,"Sheet not found")
    td=tempfile.mkdtemp(prefix="sxew_"); src=Path(td)/"model.xlsx"; shutil.copy2(MODEL,src)
    try:
        wb=openpyxl.load_workbook(src)
        ws=wb[x.sheet]
        for cell,val in x.inputs.items():
            if re.match(r"^[A-Z]{1,3}[1-9][0-9]*$",cell):
                ws[cell]=float(val) if isinstance(val,(int,float)) or str(val).replace(".","",1).isdigit() else val
        wb.save(src)
        out=Path(td)/"out"; out.mkdir()
        # LibreOffice recalculates formulas; original remains untouched.
        subprocess.run(["libreoffice","--headless","--convert-to","xlsx","--outdir",str(out),str(src)],
                       stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=120,check=True)
        result=out/"model.xlsx"
        if not result.exists(): result=src
        calcwb=openpyxl.load_workbook(result,data_only=True,read_only=True)
        wsc=calcwb[x.sheet]
        values={}
        for row in wsc.iter_rows():
            for c in row:
                if c.value is not None: values[c.coordinate]=c.value
        return {"sheet":x.sheet,"values":values,"engine":"LibreOffice-compatible Excel formula recalculation"}
    finally:
        shutil.rmtree(td,ignore_errors=True)

@app.post("/api/native-calculate")
def native_calculate(x:Calc,user=Depends(current_user)):
    try:
        return NativeWorkbook(MODEL).calculate(x.sheet,x.inputs)
    except FormulaError as e:
        raise HTTPException(422,str(e))

@app.post("/api/native-validate/{sheet}")
def native_validate(sheet:str,user=Depends(current_user)):
    # Compare the native Python engine with LibreOffice for every formula cell.
    try:
        native=NativeWorkbook(MODEL).calculate(sheet)
    except FormulaError as e:
        raise HTTPException(422,str(e))
    td=tempfile.mkdtemp(prefix="sxew_validate_"); src=Path(td)/"model.xlsx"; shutil.copy2(MODEL,src); out=Path(td)/"out";out.mkdir()
    try:
        subprocess.run(["libreoffice","--headless","--convert-to","xlsx","--outdir",str(out),str(src)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=120,check=True)
        r=out/"model.xlsx"
        wb=openpyxl.load_workbook(r,data_only=True,read_only=True); ws=wb[sheet]
        diffs=[]; checked=0
        for coord,val in native["values"].items():
            if coord not in native["formulas"]: continue
            ref=ws[coord].value; checked+=1
            if isinstance(ref,(int,float)) and isinstance(val,(int,float)):
                err=abs(ref-val); scale=max(1,abs(ref))
                if err>1e-8*scale: diffs.append({"cell":coord,"native":val,"excel":ref,"abs_error":err})
        return {"sheet":sheet,"checked":checked,"mismatches":diffs[:200],"match":not diffs}
    finally: shutil.rmtree(td,ignore_errors=True)

@app.post("/api/scenarios")
def save_scenario(x:ScenarioIn,user=Depends(current_user),s:Session=Depends(db)):
    z=Scenario(user_id=user.id,name=x.name,sheet=x.sheet,config=json.dumps(x.config,ensure_ascii=False))
    s.add(z); s.commit(); s.refresh(z); return {"id":z.id,"name":z.name}
@app.get("/api/scenarios")
def scenarios(user=Depends(current_user),s:Session=Depends(db)):
    return [{"id":z.id,"name":z.name,"sheet":z.sheet,"config":json.loads(z.config)} for z in s.scalars(select(Scenario).where(Scenario.user_id==user.id).order_by(Scenario.id.desc()))]

@app.post("/api/data")
def add_data(x:DataIn,user=Depends(current_user),s:Session=Depends(db)):
    d=Timeseries(user_id=user.id,ts=x.ts,metric=x.metric,value=x.value,period=x.period); s.add(d); s.commit(); return {"ok":True,"id":d.id}
@app.get("/api/data")
def get_data(metric:Optional[str]=None,period:Optional[str]=None,user=Depends(current_user),s:Session=Depends(db)):
    q=select(Timeseries).where(Timeseries.user_id==user.id)
    if metric:q=q.where(Timeseries.metric==metric)
    if period:q=q.where(Timeseries.period==period)
    q=q.order_by(Timeseries.ts)
    return [{"ts":d.ts.isoformat(),"metric":d.metric,"value":d.value,"period":d.period} for d in s.scalars(q)]
@app.get("/api/data.csv")
def export_csv(metric:Optional[str]=None,period:Optional[str]=None,user=Depends(current_user),s:Session=Depends(db)):
    data=get_data(metric,period,user,s); buf=io.StringIO(); w=csv.DictWriter(buf,fieldnames=["ts","metric","value","period"]);w.writeheader();w.writerows(data);buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]),media_type="text/csv",headers={"Content-Disposition":"attachment; filename=sxew-data.csv"})


class ProcessIn(BaseModel):
    circuit:dict[str,Any]

class OptimizeIn(BaseModel):
    circuit:dict[str,Any]
    variables:list[dict[str,Any]]
    objectives:list[dict[str,Any]]=[]
    max_evals:int=240
    tolerance:float=1e-8
    population_size:int=24
    seed:int=42

@app.post("/api/process/simulate")
def process_simulate(x:ProcessIn,user=Depends(current_user)):
    try: return SXEWProcessEngine(x.circuit).run()
    except ProcessError as e: raise HTTPException(422,str(e))

@app.post('/api/process/optimize')
def process_optimize(x:OptimizeIn,user=Depends(current_user)):
    try:
        return optimize(x.circuit,x.variables,x.objectives,max(10,min(1000,x.max_evals)),x.tolerance,x.population_size,x.seed)
    except OptimizationError as e: raise HTTPException(422,str(e))

@app.get('/api/process/optimization-schema')
def optimization_schema(user=Depends(current_user)):
    return {'objectives':list(__import__('app.optimizer',fromlist=['OBJECTIVES']).OBJECTIVES.keys()),'variable_examples':['nodes.E1.params.O_A','nodes.E1.params.D_Cu','nodes.E1.params.Qmax_Cu','feed.flow','feed.Cu','feed.H2SO4'],'notes':'مسیر nodes.<id>.params.<parameter> برای پارامتر واحدها و feed.<parameter> برای خوراک استفاده شود.'}

@app.get("/api/process/schema")
def process_schema(user=Depends(current_user)):
    return {"species":["Cu","Fe","H2SO4"],"units":{"flow":"user-defined mass/time","concentration":"mass fraction"},"blocks":{
        "Ex":{"ports":["raffinate","organic"],"parameters":["D_Cu","O_A","Qmax_Cu","isotherm_K","isotherm_n","acid_exponent","phase_recovery","phase_disengagement_time","fe_extraction"]},
        "Px":{"ports":["electrolyte","lean_organic"],"parameters":["D_Cu","A_O","phase_recovery","phase_disengagement_time","fe_stripping","acid_transfer"]},
        "Mixer":{"ports":["out"],"parameters":[]},
        "W":{"ports":["out"],"parameters":["fe_rejection"]},
        "EW":{"ports":["product","out"],"parameters":["cu_deposition"]}
    },"solver":{"max_iter":"integer","tolerance":"absolute component residual","relaxation":"0.05..0.99"}}

@app.get("/api/process/presets")
def process_presets(user=Depends(current_user)):
    return {"presets":["single_stage","three_stage_countercurrent"]}

@app.post("/api/process/solve")
def process_solve(x:ProcessIn,user=Depends(current_user)):
    """Explicit solver endpoint; returns convergence diagnostics, PFD stream graph and mass balance."""
    try:
        result=SXEWProcessEngine(x.circuit).run()
        return {"solver":result["solver"],"kpis":result["kpis"],"sankey":result["sankey"],"streams":result["streams"]}
    except ProcessError as e: raise HTTPException(422,str(e))

@app.get("/api/topologies")
def topologies():
    # Interpret model naming conventions such as 2Ex2Px1S as groups in series;
    # multiplicity creates parallel units within a group.
    out=[]
    pat=re.compile(r"(?:(\d+))?(Ex|Px|Wx|S)")
    for sheet in openpyxl.load_workbook(MODEL,read_only=True).sheetnames:
        m=re.match(r"(.+?)\s+(A1|A2|B1|B2|C1|C2|D1|D2|E1|E2)$",sheet)
        if not m: continue
        name=m.group(1); area=m.group(2)
        groups=[]
        for n,t in pat.findall(name):
            groups.append({"type":t,"count":int(n or 1)})
        if groups: out.append({"sheet":sheet,"area":area,"groups":groups})
    return out

@app.get("/api/health")
def health(): return {"ok":True,"time":datetime.utcnow().isoformat(),"excel_model":MODEL.name}

class CompareIn(BaseModel):
    circuits:list[dict[str,Any]]

@app.post('/api/process/compare')
def process_compare(x:CompareIn,user=Depends(current_user)):
    if not x.circuits or len(x.circuits)>12: raise HTTPException(422,'Provide 1..12 circuits')
    out=[]
    for i,c in enumerate(x.circuits,1):
        try:
            r=SXEWProcessEngine(c).run(); k=r['kpis']; s=r['solver']
            out.append({'index':i,'name':c.get('name',f'Scenario {i}'),'Cu_recovery_pct':k['Cu_recovery_pct'],'product_Cu_mass':k['product_Cu_mass'],'mass_balance_error':k['mass_balance_error'],'iterations':s['iterations'],'residual':s['residual']})
        except ProcessError as e: out.append({'index':i,'name':c.get('name',f'Scenario {i}'),'error':str(e)})
    return {'scenarios':out}

@app.post('/api/process/analysis')
def process_analysis(x:ProcessIn,user=Depends(current_user)):
    try:
        r=SXEWProcessEngine(x.circuit).run()
        return {'profiles':r.get('stage_profiles',[]),'isotherms':r.get('isotherms',[]),'kpis':r['kpis'],'solver':r['solver']}
    except ProcessError as e: raise HTTPException(422,str(e))
